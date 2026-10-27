"""Music cue sheet generation: EDL music tracks + a music-library spreadsheet -> a filled Word template.

No UI in here so it can be tested on its own.
"""
import copy
import datetime
import html
import json
import os
import re
import urllib.parse

import docx
import docx.shared
import openpyxl

import edl_report_core as core

USAGE_CODES = {
    "BI": "Background Instrumental",
    "BV": "Background Vocal",
    "VI": "Visual Instrumental",
    "VV": "Visual Vocal",
    "MT": "Main Title",
    "ET": "End Title",
    "LO": "Logo",
    "TH": "Theme",
}

INFO_FIELDS = [
    "Show Title", "Episode Title", "Cart #", "Length of Show", "Cue Sheet Type", "Production Group",
    "Show Type", "Approver", "Channel of Initial Exhibition", "Producer(s)",
    "Country/Territory of Initial Exhibition", "Executive Producer(s)", "Initial Air Date", "Created By",
    "Music Content Indicator", "Production Source Indicator", "Production Type", "Last Updated Date",
]

# header text (normalised) -> event keys to try, in order
CUE_COLUMNS = {
    "": ["@row"],  # the "#" header normalises to nothing
    "cue_num": ["cue_num", "cue_number", "cue"],
    "musical_composition": ["musical_composition", "composition", "title", "track_title", "cue_title", "@clip"],
    "performer_s": ["performer_s", "performers", "performer", "artist"],
    "writer_s": ["writer_s", "writers", "writer", "composer", "composers"],
    "publisher_s_copyright": ["publisher_s_copyright", "publisher_s", "publishers", "publisher", "copyright"],
    "time_in_time_out_time_used": ["@time"],
    "use_type": ["use_type", "usage_type", "@use"],
    "sec_use_type": ["sec_use_type", "secondary_use_type"],
    "performing_rights_society": ["performing_rights_society", "pro", "society", "societies", "rights_society"],
    "record_label_s": ["record_label_s", "record_label", "label", "labels"],
}
LIBRARY_KEYS = {"title", "composition", "musical_composition", "writer", "writers", "composer", "publisher",
                "publishers", "pro", "society", "label", "record_label"}


# --------------------------------------------------------------------------
# Library track pages (copied from the Paramount Global Music Library)
# --------------------------------------------------------------------------
LIBRARY_SEARCH_URL = "https://www.paramountglobalmusiclibrary.com/tracks?s="
PROS = ("BMI|ASCAP|SESAC|GMR|PRS|MCPS|SOCAN|APRA|AMCOS|GEMA|SACEM|SIAE|SUISA|JASRAC|SGAE|STIM|TONO|KODA|BUMA|"
        "SAYCE|SADAIC|ECAD|IMRO|TEOSTO|NCB|PRO")
ROW_TAIL_RE = re.compile(r"(?P<role>[A-Z]{1,2})\s*(?P<pro>" + PROS + r")?\s*"
                         r"(?P<ipi>\d{8,11})?\s*(?P<share>100|[1-9]\d?(?:\.\d+)?)%")
SINGLE_FIELDS = {"filename": "Filename", "artist": "Artist", "album": "Album", "library": "Library",
                 "label": "Label"}


def clean_page_text(text):
    """Pasted text, markdown or saved HTML -> plain lines (links keep their text, table cells become tabs)."""
    if re.search(r"<[a-zA-Z][^>]*>", text):
        text = re.sub(r"</t[dh]\s*>", "\t", text)
        text = re.sub(r"<br\s*/?>|</(?:tr|div|p|h\d|li|dt|dd)\s*>", "\n", text)
        text = html.unescape(re.sub(r"<[^>]+>", "", text))
    text = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", text)
    return text.replace("\r", "")


def _field(chunk, label):
    match = re.search(r"^[#*\s]*%s\s*:[*\s]*(.*)$" % re.escape(label), chunk, re.I | re.M)
    if not match:
        return ""
    if match.group(1).strip():
        return match.group(1).strip()
    rest = [line for line in chunk[match.end():].split("\n") if line.strip()]
    return rest[0].strip() if rest else ""


def _people(section):
    """'Name Role PRO IPI Share' rows -> [(name, pro, share)]."""
    section = re.sub(r"Name\s*Role\s*PRO\s*IPI\s*(?:Ownership\s*Share)?\s*(?:Collection\s*Share)?\s*(?:Share)?", "",
                     section, count=1, flags=re.I).strip()
    rows = []
    if "\t" in section:
        for line in section.split("\n"):
            cells = [c.strip() for c in line.split("\t")]
            share = next((c for c in cells[1:] if c.endswith("%")), "")
            if cells[0] and share:
                pro = cells[2] if len(cells) > 2 and not cells[2].endswith("%") else ""
                rows.append((cells[0], pro, share.rstrip("%")))
    else:
        cursor = 0
        for found in ROW_TAIL_RE.finditer(section):
            name = section[cursor:found.start()].strip(" \t\r\n[]")
            if name:
                rows.append((name, found.group("pro") or "", found.group("share")))
            cursor = found.end()
    return rows


def _share_text(rows):
    return " / ".join("{}{}{}".format(n, " ({}%)".format(sh) if sh else "", "-" + p if p else "") for n, p, sh in rows)


def _is_title(line):
    line = line.strip(" #*")
    return bool(line) and "%" not in line and ":" not in line and len(line) < 120


def parse_track_pages(text):
    """Records (dicts keyed like the cue-sheet columns) from one or more copied library track pages."""
    lines = clean_page_text(text).split("\n")
    starts = [i for i, line in enumerate(lines) if re.match(r"^[#*\s]*Length\s*:", line, re.I)]
    if not starts:
        starts = [i for i, line in enumerate(lines) if re.match(r"^[#*\s]*Filename\s*:", line, re.I)]
    if not starts:
        starts = [i for i, line in enumerate(lines) if re.match(r"^[#*\s]*Artist\s*:", line, re.I)]
    titles = []  # index of the title line above each start, if there is one
    for start in starts:
        j = start - 1
        while j >= 0 and not lines[j].strip():
            j -= 1
        titles.append(j if j >= 0 and _is_title(lines[j]) and (not titles or j > starts[len(titles) - 1]) else None)
    records = []
    for n, start in enumerate(starts):
        end = len(lines)
        if n + 1 < len(starts):
            end = titles[n + 1] if titles[n + 1] is not None else starts[n + 1]
        chunk = "\n".join(lines[start:end])
        values = {key: _field(chunk, label) for label, key in SINGLE_FIELDS.items()}
        title = lines[titles[n]].strip(" #*") if titles[n] is not None else ""
        filename = values["Filename"] or title
        if not filename:
            continue
        pubs = re.search(r"^[#*\s]*Publishers\s*:(.*?)(?=^[#*\s]*Writers\s*:|\Z)", chunk, re.I | re.M | re.S)
        writers = re.search(r"^[#*\s]*Writers\s*:(.*)\Z", chunk, re.I | re.M | re.S)
        publishers = _people(pubs.group(1)) if pubs else []
        authors = _people(writers.group(1)) if writers else []
        societies = []
        for _, pro, _ in authors + publishers:
            if pro and pro not in societies:
                societies.append(pro)
        composition = title.split(".", 1)[1].strip() if "." in title else title
        records.append({
            "Filename": filename,
            "Musical Composition": composition or filename,
            "Performer(s)": values["Artist"],
            "Writer(s)": _share_text(authors),
            "Publisher(s)": _share_text(publishers),
            "Performing Rights Society": " / ".join(societies),
            "Record Label(s)": values["Label"],
            "Library": values["Library"],
            "Album": values["Album"],
        })
    return records


def load_cache(path):
    try:
        with open(path, encoding="utf-8") as handle:
            data = json.load(handle)
        return data if isinstance(data, list) else []
    except (OSError, ValueError):
        return []


def add_to_cache(path, new_records):
    """Merge by filename (newer wins), save, and return the full list."""
    records = {core.key_name(r["Filename"]): r for r in load_cache(path)}
    for record in new_records:
        records[core.key_name(record["Filename"])] = record
    merged = list(records.values())
    os.makedirs(os.path.dirname(str(path)), exist_ok=True)
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(merged, handle, indent=2)
    return merged


def search_url(song):
    return LIBRARY_SEARCH_URL + urllib.parse.quote(str(song))


def time_cell(rec_in, rec_out):
    """'01:00:28:19','01:01:20:03' -> '01:00:28/01:01:20 (00:00:52)' (frames dropped, 1 second minimum)."""
    def seconds(tc):
        h, m, s = (int(x) for x in re.split(r"[:;.]", tc)[:3])
        return h * 3600 + m * 60 + s

    def hms(total):
        return "{:02d}:{:02d}:{:02d}".format(total // 3600, total % 3600 // 60, total % 60)

    start, end = seconds(rec_in), seconds(rec_out)
    return "{}/{} ({})".format(hms(start), hms(end), hms(max(1, end - start)))


CATALOG_RE = re.compile(r"^([A-Z]{2,5}\d{2,4}_\d{2})(?:_\d+)?[ _]+([^_]+?)(?:_.*)?$")
STEM_RE = re.compile(r"[ _-]*(?:\b(?:FULL ?MIX|STEM|MIX|ONLY)\b|\w*STEM$|FULLMIX$).*$")


def song_key(clip_name):
    """Group stems of one song: 'XRC088_08_10 THE PAST JUDGEMENT_STRINGS.WAV' -> 'XRC088_08 THE PAST JUDGEMENT'."""
    name = re.sub(r"\.(?:wav|aiff?|mp3|m4a|flac|mov|mp4)$", "", str(clip_name), flags=re.I).strip().upper()
    catalog = CATALOG_RE.match(name)
    if catalog:
        return "{} {}".format(catalog.group(1), catalog.group(2).strip())
    name = re.sub(r"[_.]+", " ", name)
    name = re.sub(r"\s+", " ", STEM_RE.sub("", name)).strip(" -")
    return name or os.path.splitext(str(clip_name))[0].upper()


def exclude_clips(events, text):
    """Drop clips whose name contains any comma-separated word, e.g. 'SFX, GFX, SIREN'."""
    words = [w.strip().upper() for w in re.split(r"[,;]", text or "") if w.strip()]
    return [e for e in events if not any(w in str(e.get("clip_name", "")).upper() for w in words)]


def merge_songs(events, fps, gap_frames):
    """One line per use of a song: stems on any track that overlap or sit within gap_frames are joined."""
    for event in events:
        event["song"] = song_key(event.get("clip_name", ""))
    ordered = sorted(events, key=lambda e: core.tc_to_frames(e["rec_in"], fps))
    open_cues, merged = {}, []
    for event in ordered:
        start = core.tc_to_frames(event["rec_in"], fps)
        cue = open_cues.get(event["song"])
        if cue and start - core.tc_to_frames(cue["rec_out"], fps) <= gap_frames:
            if core.tc_to_frames(event["rec_out"], fps) > core.tc_to_frames(cue["rec_out"], fps):
                cue["rec_out"] = event["rec_out"]
            cue["stems"] = cue.get("stems", 1) + 1
            if "FULL" in event["clip_name"].upper() and "FULL" not in cue["clip_name"].upper():
                cue["clip_name"] = event["clip_name"]
            continue
        cue = dict(event)
        open_cues[event["song"]] = cue
        merged.append(cue)
    for cue in merged:
        cue["duration_frames"] = core.tc_to_frames(cue["rec_out"], fps) - core.tc_to_frames(cue["rec_in"], fps)
    return merged


def build_cues(events, fps, track_filter="", library=None, key_column="", default_use="BI", log=None, join_seconds=2, exclude=""):
    """Turn parsed events into sorted cue dicts, one per use of a song. library: dict records matched on key_column."""
    log = log or (lambda msg: None)
    music = core.drop_gaps(core.filter_tracks(events, track_filter))
    music = exclude_clips(music, exclude)
    music = merge_songs(music, fps, int(round(join_seconds * core.nominal_fps(fps))))
    matched = 0
    # Match on the file name first, then on the song name so one library row covers every stem.
    by_song = [dict(r, **{"_song_key": song_key(r.get(key_column, ""))}) for r in library or []]
    misses = []
    for cue in music:
        if library and core.merge_metadata([cue], library, key_column)[0]:
            matched += 1
        elif library and core.merge_metadata([cue], by_song, "_song_key", event_field="song", strip_extension=False)[0]:
            matched += 1
        else:
            misses.append(cue["song"])
    for name in sorted(set(misses)):
        log("Warning: no library entry for '{}'".format(name))
    return music, matched


def event_value(event, keys, row_number):
    for key in keys:
        if key == "@row":
            return str(row_number)
        if key == "@clip":
            return str(event.get("song") or os.path.splitext(str(event.get("clip_name", "")))[0])
        if key == "@time":
            return time_cell(event["rec_in"], event["rec_out"])
        if key == "@use":
            return event.get("_default_use", "BI")
        value = event.get(key)
        if value not in (None, ""):
            return str(value)
    return ""


def _unique_cells(row):
    seen, cells = set(), []
    for cell in row.cells:
        if id(cell._tc) not in seen:
            seen.add(id(cell._tc))
            cells.append(cell)
    return cells


def set_cell_text(cell, text, size, bold=False, font="Arial"):
    paragraph = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        return
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = docx.shared.Pt(size)
    run.font.bold = bold


def find_tables(document):
    info = cues = None
    for table in document.tables:
        first = [c.text.strip() for c in table.rows[0].cells]
        if any(core.key_name(t) == "cue_num" for t in first):
            cues = table
        elif any(core.key_name(t) == "show_title" for t in first) or info is None:
            info = table
    if cues is None:
        raise ValueError("The template has no cue table (a table with a 'Cue Num' header).")
    return info, cues


def fill_info_table(table, values):
    """Write each value in the cell to the right of its label (labels matched ignoring punctuation)."""
    wanted = {core.key_name(k): v for k, v in values.items()}
    for row in table.rows:
        cells = _unique_cells(row)
        for i in range(0, len(cells) - 1):
            key = core.key_name(cells[i].text)
            if key in wanted and not cells[i + 1].text.strip():
                set_cell_text(cells[i + 1], str(wanted[key]), 10.5)


def usage_legend(codes):
    codes = [c for c in USAGE_CODES if c in codes] + sorted(c for c in codes if c not in USAGE_CODES)
    return ", ".join("{}={}".format(c, USAGE_CODES.get(c, c)) for c in codes)


def fill_paragraphs(document, codes, created_by, sheet_id, now):
    for paragraph in document.paragraphs:
        text = paragraph.text
        if text.startswith("Usage Type:") and len(paragraph.runs) >= 2 and codes:
            lead = paragraph.runs[1].text[:len(paragraph.runs[1].text) - len(paragraph.runs[1].text.lstrip())]
            paragraph.runs[1].text = lead + usage_legend(codes)
            for run in paragraph.runs[2:]:
                run.text = ""
        elif text.startswith("*Report Generated on"):
            stamp = "{} {}".format(now.strftime("%m/%d/%Y"), now.strftime("%I:%M %p").lstrip("0"))
            runs = paragraph.runs
            for run in runs:
                if run.text.startswith("*Report Generated on"):
                    run.text = "*Report Generated on {}{}".format(stamp, " by " + created_by if created_by else "")
                elif run.text.startswith("Cue sheet ID"):
                    run.text = "Cue sheet ID: {}".format(sheet_id)


def fill_cue_table(table, cues, default_use="BI"):
    headers = [core.key_name(c.text) for c in _unique_cells(table.rows[0])]
    prototype = table.rows[1]._tr
    for number, event in enumerate(cues, 1):
        event.setdefault("_default_use", default_use)
        table._tbl.append(copy.deepcopy(prototype))
        for cell, header in zip(_unique_cells(table.rows[-1]), headers):
            keys = CUE_COLUMNS.get(header, [header])
            set_cell_text(cell, event_value(event, keys, number), 8)
    table._tbl.remove(prototype)


def program_length(events, fps):
    ins = [core.tc_to_frames(e["rec_in"], fps) for e in events if e.get("rec_in")]
    outs = [core.tc_to_frames(e["rec_out"], fps) for e in events if e.get("rec_out")]
    if not ins or not outs:
        return ""
    seconds = round((max(outs) - min(ins)) / core.nominal_fps(fps))
    return "{:02d}:{:02d}:{:02d}".format(seconds // 3600, seconds % 3600 // 60, seconds % 60)


def build_cue_sheet(source_path, template_path, output_path, *, fps=core.DEFAULT_FPS, track_filter="",
                    library_path="", key_column="", info=None, sheet_id="", default_use="BI",
                    join_seconds=2, exclude="", cache_records=None, missing=None, log=None, now=None):
    """Parse the EDL, fill the Word template, save. Returns the number of cues written."""
    log = log or (lambda msg: None)
    now = now or datetime.datetime.now()
    info = dict(info or {})
    log("Starting: {}".format(os.path.basename(source_path)))
    events = core.parse_file(source_path, fps)[0]
    log("Found {} events".format(len(events)))
    library = []
    if library_path:
        headers, rows = core.read_table(library_path)
        key_column = key_column or (headers[0] if headers else "")
        log("Found {} library entries (matching on '{}')".format(len(rows), key_column))
        library = [dict(r, Filename=r.get(key_column, "")) for r in rows]
    if cache_records:
        log("Found {} saved library tracks".format(len(cache_records)))
        library = library + list(cache_records)
    if library:
        key_column = "Filename"
    warnings = []
    cues, matched = build_cues(events, fps, track_filter, library, key_column, default_use,
                               lambda m: (warnings.append(m), log(m)), join_seconds, exclude)
    if missing is not None:
        missing.extend(re.findall(r"no library entry for '(.*)'", "\n".join(warnings)))
    if not cues:
        raise ValueError("No music events found; check the music tracks field.")
    document = docx.Document(template_path)
    info_table, cue_table = find_tables(document)
    if not info.get("Length of Show") and not info.get("length_of_show"):
        info["Length of Show"] = program_length(events, fps)
        log("Length of Show filled from the EDL: {}".format(info["Length of Show"]))
    info.setdefault("Last Updated Date", now.strftime("%m/%d/%Y"))
    if info_table is not None:
        fill_info_table(info_table, info)
    fill_cue_table(cue_table, cues, default_use)
    codes = {event_value(e, ["use_type", "usage_type", "@use"], 0) for e in cues}
    fill_paragraphs(document, codes, info.get("Created By", ""), sheet_id, now)
    document.save(output_path)
    log("Processed {} cues ({} matched the library)".format(len(cues), matched))
    log("Done: {}".format(output_path))
    return len(cues)


def write_library_template(source_path, output_path, fps=core.DEFAULT_FPS, track_filter="", exclude=""):
    """A spreadsheet of the music clips in the EDL, ready to fill with writers, publishers, etc."""
    events = core.parse_file(source_path, fps)[0]
    music = exclude_clips(core.drop_gaps(core.filter_tracks(events, track_filter)), exclude)
    names = sorted({song_key(e["clip_name"]) for e in music})
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Song", "Title", "Performer(s)", "Writer(s)", "Publisher(s)", "Performing Rights Society",
                  "Record Label(s)", "Use Type"])
    for name in names:
        sheet.append([name, name.title()])
    workbook.save(output_path)
    return len(names)


def track_summary(events):
    """[(track, count, sample clip)] for choosing the music tracks."""
    tracks = {}
    for event in events:
        entry = tracks.setdefault(event.get("track", ""), [0, ""])
        entry[0] += 1
        if not entry[1] and event.get("clip_name"):
            entry[1] = event["clip_name"]
    return [(t, c, s) for t, (c, s) in tracks.items()]
