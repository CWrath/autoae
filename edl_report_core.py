"""Core logic for EDL Reports: parse EDL / ALE / XML, merge cue metadata, write styled Excel.

No UI in here so it can be tested on its own. Every event is a dict with these keys
(plus any extra "KEY: value" comments or ALE columns, lower-cased with underscores):

    event, reel, track, edit, clip_name, source_file, comments,
    rec_in, rec_out, src_in, src_out, duration_frames
"""

import csv
import fnmatch
import json
import os
import re
import xml.etree.ElementTree as ET
from datetime import datetime
from fractions import Fraction
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

TC_RE = re.compile(r"^\d{1,2}[:;]\d{2}[:;]\d{2}[:;.]\d{2,3}$")
DEFAULT_FPS = 24.0


# --------------------------------------------------------------------------
# Timecode
# --------------------------------------------------------------------------
def nominal_fps(fps):
    return int(round(float(fps)))


def is_drop_rate(fps):
    return abs(float(fps) - 29.97) < 0.01 or abs(float(fps) - 59.94) < 0.01


def tc_to_frames(tc, fps, drop=None):
    """HH:MM:SS:FF (';' before the frames means drop frame) -> frame count."""
    tc = str(tc).strip()
    if drop is None:
        drop = ";" in tc and is_drop_rate(fps)
    h, m, s, f = (int(part) for part in re.split(r"[:;.]", tc))
    base = nominal_fps(fps)
    frames = ((h * 60 + m) * 60 + s) * base + f
    if drop:
        skip = 2 if base == 30 else 4
        total_minutes = h * 60 + m
        frames -= skip * (total_minutes - total_minutes // 10)
    return frames


def frames_to_tc(frames, fps, drop=False):
    frames = max(0, int(frames))
    base = nominal_fps(fps)
    sep = ":"
    if drop and is_drop_rate(fps):
        skip = 2 if base == 30 else 4
        per_10 = base * 600 - skip * 9
        per_min = base * 60 - skip
        tens, rest = divmod(frames, per_10)
        if rest > skip:
            frames += skip * 9 * tens + skip * ((rest - skip) // per_min)
        else:
            frames += skip * 9 * tens
        sep = ";"
    f = frames % base
    s = (frames // base) % 60
    m = (frames // (base * 60)) % 60
    h = frames // (base * 3600)
    return "{:02d}:{:02d}:{:02d}{}{:02d}".format(h, m, s, sep, f)


def format_duration(frames, fps, style="tc"):
    seconds = frames / float(fps) if fps else 0
    if style == "frames":
        return int(frames)
    if style == "seconds":
        return round(seconds, 2)
    if style == "mmss":
        whole = int(round(seconds))
        return "{}:{:02d}".format(whole // 60, whole % 60)
    return frames_to_tc(frames, fps)  # duration as timecode, never drop-frame flagged


def key_name(text):
    """Normalise a column / token name: 'Source File' -> 'source_file'."""
    return re.sub(r"[^a-z0-9]+", "_", str(text).strip().lower()).strip("_")


# --------------------------------------------------------------------------
# Parsers
# --------------------------------------------------------------------------
def read_text(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("latin-1", errors="replace")


def detect_type(path):
    suffix = Path(path).suffix.lower()
    if suffix in (".edl", ".ale", ".xml", ".fcpxml"):
        return "xml" if suffix == ".fcpxml" else suffix[1:]
    head = read_text(path)[:2000].lower()
    if "heading" in head and "column" in head:
        return "ale"
    if head.lstrip().startswith("<?xml") or head.lstrip().startswith("<"):
        return "xml"
    return "edl"


def _finish(event, fps):
    """Fill duration_frames and make sure every standard key exists."""
    for key in ("event", "reel", "track", "edit", "clip_name", "source_file", "comments",
                "rec_in", "rec_out", "src_in", "src_out"):
        event.setdefault(key, "")
    if event.get("rec_in") and event.get("rec_out"):
        event["duration_frames"] = max(0, tc_to_frames(event["rec_out"], fps) - tc_to_frames(event["rec_in"], fps))
    elif event.get("src_in") and event.get("src_out"):
        event["duration_frames"] = max(0, tc_to_frames(event["src_out"], fps) - tc_to_frames(event["src_in"], fps))
    else:
        event.setdefault("duration_frames", 0)
    return event


def parse_edl(path, fps=DEFAULT_FPS):
    events = []
    title = ""
    current = None
    for line in read_text(path).splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.upper().startswith("TITLE:"):
            title = stripped[6:].strip()
            continue
        if stripped.upper().startswith("FCM:"):
            continue
        tokens = stripped.split()
        if tokens[0].isdigit() and len(tokens) >= 8 and all(TC_RE.match(t) for t in tokens[-4:]):
            current = {
                "event": int(tokens[0]), "reel": tokens[1], "track": tokens[2], "edit": tokens[3],
                "src_in": tokens[-4], "src_out": tokens[-3], "rec_in": tokens[-2], "rec_out": tokens[-1],
            }
            events.append(current)
        elif stripped.startswith("*") and current is not None:
            comment = stripped.lstrip("*").strip()
            current["comments"] = (current.get("comments", "") + " | " + comment).strip(" |")
            if ":" in comment:
                label, _, value = comment.partition(":")
                label, value = label.strip().upper(), value.strip()
                if label == "FROM CLIP NAME":
                    current["clip_name"] = value
                elif label == "TO CLIP NAME":
                    current.setdefault("to_clip_name", value)
                elif label == "SOURCE FILE":
                    current["source_file"] = value
                elif label == "LOC":
                    parts = value.split(None, 2)
                    current["marker"] = (parts[2] if len(parts) > 2 else value)
                else:
                    current.setdefault(key_name(label), value)
    for event in events:
        if not event.get("clip_name"):
            # Avid pull lists cut the reel off at 32 characters; *SOURCE FILE has the full name
            event["clip_name"] = event.get("to_clip_name") or event.get("source_file") or event["reel"]
        _finish(event, fps)
    return events, {"title": title, "fps": fps}


ALE_MAP = {"name": "clip_name", "tape": "reel", "source_file": "source_file", "tracks": "track",
           "start": "src_in", "end": "src_out", "duration": "ale_duration", "comments": "comments"}


def parse_ale(path, fps=DEFAULT_FPS):
    lines = read_text(path).splitlines()
    ale_fps = fps
    columns, rows, section = [], [], ""
    for line in lines:
        stripped = line.strip("\r\n")
        low = stripped.strip().lower()
        if low in ("heading", "column", "data"):
            section = low
            continue
        if not stripped.strip():
            continue
        if section == "heading":
            parts = stripped.split("\t")
            if parts[0].strip().upper() == "FPS" and len(parts) > 1:
                try:
                    ale_fps = float(parts[1])
                except ValueError:
                    pass
        elif section == "column" and not columns:
            columns = [c.strip() for c in stripped.split("\t")]
        elif section == "data":
            rows.append(stripped.split("\t"))
    if not columns:
        raise ValueError("No Column section found; is this an Avid Log Exchange (.ale) file?")
    events = []
    for number, row in enumerate(rows, 1):
        event = {"event": number}
        for column, value in zip(columns, row):
            value = value.strip()
            key = key_name(column)
            event[ALE_MAP.get(key, key)] = value
        if event.get("src_in"):
            event["rec_in"], event["rec_out"] = event.get("src_in"), event.get("src_out", "")
        events.append(_finish(event, ale_fps))
    return events, {"title": Path(path).stem, "fps": ale_fps}


def _local(tag):
    return tag.rsplit("}", 1)[-1]


def _text(node, path):
    found = node.find(path)
    return (found.text or "").strip() if found is not None and found.text else ""


def parse_xml(path, fps=DEFAULT_FPS):
    root = ET.fromstring(read_text(path).encode("utf-8"))
    kind = _local(root.tag)
    if kind == "xmeml":
        return _parse_xmeml(root, fps)
    if kind == "fcpxml":
        return _parse_fcpxml(root, fps)
    raise ValueError("Unsupported XML ({}). Export Final Cut Pro 7 XML or FCPXML.".format(kind))


def _parse_xmeml(root, fps):
    sequence = root.find(".//sequence")
    if sequence is None:
        raise ValueError("No sequence found in this XML.")
    timebase = _text(sequence, "rate/timebase")
    if timebase:
        fps = float(timebase) * (1000 / 1001 if _text(sequence, "rate/ntsc").upper() == "TRUE" else 1)
    drop = _text(sequence, "timecode/displayformat").upper() == "DF"
    offset = 0
    start_string = _text(sequence, "timecode/string")
    if start_string:
        offset = tc_to_frames(start_string, fps)
    files = {}
    for node in sequence.iter("file"):
        file_id = node.get("id")
        if file_id and (node.find("name") is not None):
            files.setdefault(file_id, (_text(node, "name"), _text(node, "pathurl")))
    events = []
    for media, label in (("video", "V"), ("audio", "A")):
        container = sequence.find("media/" + media)
        if container is None:
            continue
        for index, track in enumerate(container.findall("track"), 1):
            for item in track.findall("clipitem"):
                try:
                    start, end = int(_text(item, "start")), int(_text(item, "end"))
                except ValueError:
                    continue
                if start < 0 or end < 0:
                    continue
                file_node = item.find("file")
                file_id = file_node.get("id") if file_node is not None else None
                name, url = files.get(file_id, ("", ""))
                clip = _text(item, "name") or name
                events.append(_finish({
                    "track": "{}{}".format(label, index), "clip_name": clip, "reel": name or clip,
                    "source_file": url, "edit": "C",
                    "rec_in": frames_to_tc(start + offset, fps, drop), "rec_out": frames_to_tc(end + offset, fps, drop),
                    "src_in": frames_to_tc(int(_text(item, "in") or 0), fps),
                    "src_out": frames_to_tc(int(_text(item, "out") or 0), fps),
                    "comments": _text(item, "comments/mastercomment1"),
                }, fps))
    events.sort(key=lambda e: tc_to_frames(e["rec_in"], fps))
    for number, event in enumerate(events, 1):
        event["event"] = number
    return events, {"title": _text(sequence, "name"), "fps": fps}


def _rational(value):
    """FCPXML time such as '1001/30000s' or '5s' -> Fraction of seconds."""
    if not value:
        return Fraction(0)
    value = value.rstrip("s")
    return Fraction(value) if "/" in value else Fraction(int(float(value) * 1000), 1000)


def _parse_fcpxml(root, fps):
    fmt = root.find(".//format")
    if fmt is not None and fmt.get("frameDuration"):
        fps = float(1 / _rational(fmt.get("frameDuration")))
    drop = False
    sequence = root.find(".//sequence")
    if sequence is not None:
        drop = sequence.get("tcFormat", "NDF") == "DF"
    start_offset = _rational(sequence.get("tcStart")) if sequence is not None else Fraction(0)
    clip_tags = {"clip", "asset-clip", "ref-clip", "sync-clip", "audio", "video", "mc-clip"}
    events = []

    def walk(node, parent_start, parent_record, lane_label):
        for child in node:
            tag = _local(child.tag)
            if tag in clip_tags:
                offset = _rational(child.get("offset"))
                record = parent_record + (offset - parent_start)
                duration = _rational(child.get("duration"))
                lane = child.get("lane")
                label = "A{}".format(lane) if lane and int(lane) > 0 else ("V" if not lane or int(lane) == 0 else "B")
                events.append(_finish({
                    "track": label, "clip_name": child.get("name", ""), "reel": child.get("name", ""),
                    "edit": "C", "source_file": "",
                    "rec_in": frames_to_tc(int((record + start_offset) * fps), fps, drop),
                    "rec_out": frames_to_tc(int((record + duration + start_offset) * fps), fps, drop),
                    "src_in": "", "src_out": "",
                    "comments": child.get("note", "") or "",
                }, fps))
                walk(child, _rational(child.get("start")), record, label)
            elif tag in ("spine", "gap", "sequence", "project", "event"):
                if tag == "gap":
                    walk(child, _rational(child.get("start")), parent_record + (_rational(child.get("offset")) - parent_start), "V")
                else:
                    walk(child, parent_start, parent_record, lane_label)

    walk(root, Fraction(0), Fraction(0), "V")
    events.sort(key=lambda e: tc_to_frames(e["rec_in"], fps))
    for number, event in enumerate(events, 1):
        event["event"] = number
    title = ""
    if root.find(".//project") is not None:
        title = root.find(".//project").get("name", "")
    return events, {"title": title, "fps": fps}


def parse_file(path, fps=DEFAULT_FPS):
    kind = detect_type(path)
    if kind == "ale":
        return parse_ale(path, fps)
    if kind == "xml":
        return parse_xml(path, fps)
    return parse_edl(path, fps)


# --------------------------------------------------------------------------
# Processing
# --------------------------------------------------------------------------
def expand_track_patterns(text):
    """'A13-A17, V*' -> ['a13', 'a14', 'a15', 'a16', 'a17', 'v*']."""
    wanted = []
    for part in re.split(r"[,;]", text or ""):
        part = part.strip().lower()
        if not part:
            continue
        span = re.fullmatch(r"([a-z]*)(\d+)\s*-\s*(?:\1)?(\d+)", part)
        if span and int(span.group(2)) <= int(span.group(3)):
            wanted.extend("{}{}".format(span.group(1), n) for n in range(int(span.group(2)), int(span.group(3)) + 1))
        else:
            wanted.append(part)
    return wanted


def filter_tracks(events, patterns):
    """patterns: 'A, A2, A13-A17' or 'A*'; blank keeps everything. Matches track names, case-insensitive."""
    wanted = expand_track_patterns(patterns)
    if not wanted:
        return list(events)
    return [e for e in events
            if any(fnmatch.fnmatch(str(e.get("track", "")).lower(), pattern) for pattern in wanted)]


def drop_gaps(events):
    """Remove black / slug / filler lines that carry no clip of their own."""
    skip = ("BL", "BLACK", "SLUG", "GAP")
    return [e for e in events if str(e.get("reel", "")).upper() not in skip and e.get("clip_name")]


def collapse_duplicates(events):
    """Stereo pairs and linked tracks (A, A2) are one use of a clip: keep one line, list the tracks."""
    seen, result = {}, []
    for event in events:
        key = (event.get("clip_name"), event.get("rec_in"), event.get("rec_out"))
        if key in seen:
            tracks = seen[key]["track"].split("/")
            if event.get("track") and event["track"] not in tracks:
                seen[key]["track"] = "/".join(tracks + [event["track"]])
            continue
        copy = dict(event)
        seen[key] = copy
        result.append(copy)
    return result


def merge_back_to_back(events, fps, gap_frames=0):
    """Join consecutive events of the same clip on the same track into one cue.

    A single music cue is often cut into several EDL events; reports want it once.
    """
    ordered = sorted(events, key=lambda e: (str(e.get("track", "")), tc_to_frames(e["rec_in"], fps)))
    merged = []
    for event in ordered:
        last = merged[-1] if merged else None
        if (last and last["clip_name"] == event["clip_name"] and last["track"] == event["track"]
                and tc_to_frames(event["rec_in"], fps) - tc_to_frames(last["rec_out"], fps) <= gap_frames):
            last["rec_out"] = event["rec_out"]
            last["src_out"] = event.get("src_out", last.get("src_out", ""))
            last["duration_frames"] = (tc_to_frames(last["rec_out"], fps) - tc_to_frames(last["rec_in"], fps))
            last["merged_events"] = last.get("merged_events", 1) + 1
            continue
        merged.append(dict(event))
    merged.sort(key=lambda e: (tc_to_frames(e["rec_in"], fps), str(e.get("track", ""))))
    return merged


def exact_fps(fps):
    """23.976 -> 24000/1001 and so on, so conversions between broadcast rates don't drift."""
    fps = float(fps)
    for nominal in (24, 30, 60):
        if abs(fps - nominal * 1000 / 1001) < 0.01:
            return Fraction(nominal * 1000, 1001)
    return Fraction(fps).limit_denominator(1000)


def to_df_timecode(frames, fps):
    """A running time in `fps` frames -> the same real time as 29.97 drop-frame, written 00;07;50;00."""
    df_frames = int(frames * Fraction(30000, 1001) / exact_fps(fps) + Fraction(1, 2))
    return frames_to_tc(df_frames, 29.97, drop=True).replace(":", ";")


def round_df_timecode(df_text):
    """Round 00;06;15;29 to the nearest whole second: 00;06;16;00."""
    h, m, s, f = (int(part) for part in df_text.split(";"))
    total = (h * 60 + m) * 60 + s + (1 if f >= 15 else 0)
    return "{:02d};{:02d};{:02d};00".format(total // 3600, (total // 60) % 60, total % 60)


def _is_black(event):
    reel = str(event.get("reel", "")).strip().upper()
    name = str(event.get("clip_name", "")).strip().upper()
    return reel in ("BL", "BLACK", "SLUG", "GAP") or name in ("BL", "BLACK", "SLUG", "GAP")


def build_segments(events, fps, min_break_frames=1, max_break_frames=None):
    """Split a cut into program segments and the breaks between them.

    Black / slug events and empty gaps are breaks; everything else is program.
    If the EDL has video tracks only those are used, so audio under a break can't hide it.
    Returns [{"kind": "segment"|"break", "label", "start", "end", "frames", "black"}],
    where end is exclusive. Leading and trailing breaks are dropped.
    """
    video = [e for e in events if str(e.get("track", "")).upper().startswith(("V", "B"))]
    pool = [e for e in (video or events) if e.get("rec_in") and e.get("rec_out")]
    spans = []
    for event in pool:
        start, end = tc_to_frames(event["rec_in"], fps), tc_to_frames(event["rec_out"], fps)
        if end > start:
            spans.append((start, end, _is_black(event)))
    if not spans:
        return []
    points = sorted({p for s, e, _ in spans for p in (s, e)})
    pieces = []  # (start, end, kind)
    for a, b in zip(points, points[1:]):
        covering = [black for s, e, black in spans if s <= a and e >= b]
        if not covering:
            kind = "gap"
        elif all(covering):
            kind = "black"
        else:
            kind = "segment"
        if pieces and pieces[-1][2] == kind and pieces[-1][1] == a:
            pieces[-1] = (pieces[-1][0], b, kind)
        else:
            pieces.append((a, b, kind))
    while pieces and pieces[0][2] != "segment":
        pieces.pop(0)
    while pieces and pieces[-1][2] != "segment":
        pieces.pop()

    # a gap outside the break length range is black inside the program, not a break
    folded = []
    for start, end, kind in pieces:
        length = end - start
        if kind != "segment" and (length < min_break_frames or (max_break_frames and length > max_break_frames)):
            kind = "segment"
        if folded and folded[-1][2] == "segment" and kind == "segment":
            folded[-1] = (folded[-1][0], end, "segment")
        else:
            folded.append((start, end, kind))

    result, seg_count, break_count = [], 0, 0
    for start, end, kind in folded:
        if kind == "segment":
            seg_count += 1
            label = "SEGMENT {}".format(seg_count)
        else:
            break_count += 1
            label = "Break {} (Black)".format(break_count) if kind == "black" else "Break {}".format(break_count)
        result.append({"kind": "segment" if kind == "segment" else "break", "label": label,
                       "number": seg_count if kind == "segment" else break_count,
                       "start": start, "end": end, "frames": end - start, "black": kind == "black"})
    return result


def write_segment_sheet(segments, preset, ctx, output_path):
    """Segment timing sheet: in / out, length, and the length converted to 29.97 drop frame."""
    style = preset.get("style", {})
    fps = ctx["fps"]
    font_name = style.get("font", "Arial")
    size = style.get("size", 11)
    head_size = style.get("header_size", 10)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\\/*?:\[\]]", "", preset.get("sheet", "Timing"))[:31] or "Timing"

    row = 1
    for line in preset.get("info", []):
        sheet.cell(row, 1, _fill_text(line, ctx)).font = Font(name=font_name, size=head_size, bold=True)
        row += 1
    row += 1

    label = "{:.2f}".format(float(fps))
    headings = ["{} Timecode In".format(label), "{} Timecode Out".format(label), "", "Length",
                "DF Conversion", "DF Rounded"]
    for index, text in enumerate(headings, 1):
        cell = sheet.cell(row, index, text)
        cell.font = Font(name=font_name, size=head_size, bold=True)
        cell.alignment = center
        cell.border = border
    for letter, width in zip("ABCDEF", (20, 20, 20, 16, 16, 16)):
        sheet.column_dimensions[letter].width = width
    header_row = row
    row += 1

    previous = None
    for item in segments:
        if previous and previous["kind"] == "break":
            row += 1  # blank row between a break and the next segment, as in the delivered sheets
        label_text = item["label"]
        if item["kind"] == "break" and preset.get("break_label"):
            label_text = preset["break_label"].format(n=item["number"])
        values = [frames_to_tc(item["start"], fps), frames_to_tc(item["end"] - 1, fps), label_text,
                  frames_to_tc(item["frames"], fps), None, None]
        df = to_df_timecode(item["frames"], fps)
        values[4] = df
        if item["kind"] == "segment":
            values[5] = round_df_timecode(df)
        for index, value in enumerate(values, 1):
            cell = sheet.cell(row, index, value)
            bold = index == 3
            cell.font = Font(name=font_name, size=head_size if bold else size, bold=bold)
            cell.alignment = center
            if value is not None:
                cell.border = border
        previous = item
        row += 1

    program = [s for s in segments if s["kind"] == "segment"]
    if program:
        if preset.get("tpt", "span") == "span":
            total = (program[-1]["end"] - 1) - program[0]["start"]
        else:
            total = sum(s["frames"] for s in program)
        row += 1
        # The DF total is program time only (no breaks) less any card that isn't counted, e.g. a credits card
        df_total = sum(s["frames"] for s in program) - int(round(float(ctx.get("df_deduct_seconds") or 0) * fps))
        entries = [(3, "TPT in {}".format(label)), (4, frames_to_tc(total, fps)),
                   (5, "TPT in 29.97/DF"), (6, to_df_timecode(max(0, df_total), fps).replace(";", ":"))]
        for column, value in entries:
            cell = sheet.cell(row, column, value)
            cell.font = Font(name=font_name, size=head_size if column in (3, 5) else size, bold=column in (3, 5))
            cell.alignment = center
            cell.border = border
        sheet.column_dimensions["E"].width = 20
        sheet.column_dimensions["F"].width = 18
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.print_title_rows = "{0}:{0}".format(header_row)
    workbook.save(output_path)
    return output_path


def _norm_match(value, strip_extension):
    text = str(value or "").strip().casefold()
    if strip_extension:
        text = os.path.splitext(text)[0]
    return re.sub(r"\s+", " ", text)


def read_table(path):
    """First sheet of an .xlsx or a .csv as a list of dicts keyed by the header row."""
    path = str(path)
    if path.lower().endswith((".xlsx", ".xlsm")):
        sheet = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    else:
        rows = list(csv.reader(read_text(path).splitlines()))
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return [], []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    records = [{headers[i]: ("" if v is None else v) for i, v in enumerate(r) if i < len(headers) and headers[i]}
               for r in rows[1:]]
    return headers, records


def merge_metadata(events, records, key_column, event_field="clip_name", strip_extension=True):
    """Add the spreadsheet's columns to events whose clip name matches the key column.

    Returns (matched_count, unmatched_event_names, unused_record_keys).
    """
    lookup = {}
    for record in records:
        lookup.setdefault(_norm_match(record.get(key_column, ""), strip_extension), record)
    used, unmatched = set(), []
    for event in events:
        key = _norm_match(event.get(event_field, ""), strip_extension)
        record = lookup.get(key)
        if record is None:
            unmatched.append(event.get(event_field, ""))
            continue
        used.add(key)
        for column, value in record.items():
            name = key_name(column)
            if name and event.get(name) in (None, ""):
                event[name] = value
    unused = [k for k in lookup if k and k not in used]
    return len(events) - len(unmatched), sorted(set(unmatched)), unused


# --------------------------------------------------------------------------
# Presets and field resolution
# --------------------------------------------------------------------------
BUILTIN_PRESETS = {
    "Clearance Report": {
        "sheet": "Clearance",
        "title": "{project}  |  Clearance Report",
        "info": ["EDL: {source_name}", "Generated: {generated}"],
        "sort": "rec_in",
        "merge_back_to_back": True,
        "columns": [
            {"header": "#", "field": "row", "width": 5},
            {"header": "Cue / Clip", "field": "clip_name", "width": 38, "aliases": ["Title", "Cue Title", "Clip Name"]},
            {"header": "Source File", "field": "source_file", "width": 30},
            {"header": "TC In", "field": "rec_in", "width": 13, "aliases": ["Start", "Record In"]},
            {"header": "TC Out", "field": "rec_out", "width": 13, "aliases": ["End", "Record Out"]},
            {"header": "Duration", "field": "duration", "width": 13, "aliases": ["Length", "Dur"]},
            {"header": "Track", "field": "track", "width": 8},
            {"header": "Licensor", "field": "licensor", "width": 24, "aliases": ["Owner", "Rights Holder"]},
            {"header": "Territory", "field": "territory", "width": 14},
            {"header": "Term", "field": "term", "width": 14},
            {"header": "Clearance Status", "field": "status", "width": 18, "aliases": ["Status", "Cleared"]},
            {"header": "Notes", "field": "{notes}", "width": 34, "aliases": ["Comments", "Remarks"]},
        ],
        "totals": "Total runtime",
        "style": {"font": "Calibri", "size": 11, "header_fill": "1F3864", "header_font": "FFFFFF",
                  "zebra": "EEF1F7", "freeze": True, "autofilter": True, "landscape": True},
    },
    "Music Cue Sheet": {
        "sheet": "Cue Sheet",
        "title": "{project}  |  Music Cue Sheet",
        "info": ["EDL: {source_name}", "Generated: {generated}"],
        "sort": "rec_in",
        "merge_back_to_back": True,
        "columns": [
            {"header": "Cue #", "field": "row", "width": 7},
            {"header": "Title", "field": "title|clip_name", "width": 36},
            {"header": "Composer", "field": "composer", "width": 24},
            {"header": "Publisher", "field": "publisher", "width": 24},
            {"header": "Label", "field": "label", "width": 20},
            {"header": "Usage", "field": "usage", "width": 14},
            {"header": "TC In", "field": "rec_in", "width": 13},
            {"header": "TC Out", "field": "rec_out", "width": 13},
            {"header": "Duration", "field": "duration_mmss", "width": 10},
        ],
        "totals": "Total music",
        "style": {"font": "Calibri", "size": 11, "header_fill": "2E4A3F", "header_font": "FFFFFF",
                  "zebra": "F0F5F2", "freeze": True, "autofilter": True, "landscape": True},
    },
    "Segment Timing Sheet": {
        "layout": "segments",
        "sheet": "Timing Sheet",
        "info": ["Series Title: {series}", "Episode Title: {episode_title}",
                 "Episode Number: {episode_number}", "Program ID: {program_id}"],
        "tpt": "span",
        "df_deduct_seconds": 2,
        "break_label": "Break {n} (Black)",
        "break_seconds": 2,
        "style": {"font": "Arial", "size": 11, "header_size": 10},
    },
    "Timing Sheet": {
        "sheet": "Timing",
        "title": "{project}  |  Timing Sheet",
        "info": ["EDL: {source_name}", "Generated: {generated}"],
        "sort": "rec_in",
        "merge_back_to_back": False,
        "columns": [
            {"header": "Event", "field": "event", "width": 8},
            {"header": "Clip", "field": "clip_name", "width": 40},
            {"header": "Reel", "field": "reel", "width": 16},
            {"header": "Track", "field": "track", "width": 8},
            {"header": "Record In", "field": "rec_in", "width": 14},
            {"header": "Record Out", "field": "rec_out", "width": 14},
            {"header": "Duration", "field": "duration", "width": 14},
            {"header": "Runtime", "field": "duration_mmss", "width": 10},
        ],
        "totals": "Total",
        "style": {"font": "Calibri", "size": 11, "header_fill": "333333", "header_font": "FFFFFF",
                  "zebra": "F4F4F4", "freeze": True, "autofilter": True, "landscape": True},
    },
}


def load_user_presets(path):
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return {name: value for name, value in data.items()
            if isinstance(name, str) and isinstance(value, dict) and isinstance(value.get("columns"), list)}


def all_presets(user_path):
    presets = {name: dict(value) for name, value in BUILTIN_PRESETS.items()}
    presets.update(load_user_presets(user_path))
    return presets


def make_context(project, source_path, meta, fields=None):
    context = {"series": "", "episode_title": "", "episode_number": "", "program_id": "N/A"}
    context.update({key_name(k): v for k, v in (fields or {}).items() if v not in (None, "")})
    return {**context, "project": project or meta.get("title") or Path(source_path).stem,
            "source_name": Path(source_path).name, "edl_title": meta.get("title", ""),
            "generated": datetime.now().strftime("%Y-%m-%d %H:%M"), "fps": meta.get("fps", DEFAULT_FPS)}


def resolve(spec, event, row_number, ctx):
    """A column 'field' is a name ('rec_in'), a fallback chain ('title|clip_name'), or a template ('{a} - {b}')."""
    fps = ctx["fps"]
    frames = event.get("duration_frames", 0)
    virtual = {"row": row_number, "duration": format_duration(frames, fps, "tc"),
               "duration_frames": frames, "duration_seconds": format_duration(frames, fps, "seconds"),
               "duration_mmss": format_duration(frames, fps, "mmss")}

    def lookup(name):
        name = key_name(name)
        if name in virtual:
            return virtual[name]
        if name in ctx:
            return ctx[name]
        value = event.get(name, "")
        return "" if value is None else value

    if "{" in spec:
        return re.sub(r"\{([^{}]+)\}", lambda m: str(lookup(m.group(1))), spec).strip()
    for name in spec.split("|"):
        value = lookup(name)
        if value not in ("", None):
            return value
    return ""


def build_rows(events, preset, ctx):
    rows = []
    for number, event in enumerate(events, 1):
        rows.append([resolve(col["field"], event, number, ctx) for col in preset["columns"]])
    return rows


def _fill_text(text, ctx):
    return re.sub(r"\{([^{}]+)\}", lambda m: str(ctx.get(key_name(m.group(1)), m.group(0))), text)


# --------------------------------------------------------------------------
# Writers
# --------------------------------------------------------------------------
def write_xlsx(events, preset, ctx, output_path):
    """Write a styled workbook from a preset and return the output path."""
    style = preset.get("style", {})
    columns = preset["columns"]
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[\\/*?:\[\]]", "", preset.get("sheet", "Report"))[:31] or "Report"
    font_name, size = style.get("font", "Calibri"), style.get("size", 11)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row = 1
    if preset.get("title"):
        sheet.cell(row, 1, _fill_text(preset["title"], ctx)).font = Font(name=font_name, size=size + 5, bold=True)
        row += 1
    for line in preset.get("info", []):
        sheet.cell(row, 1, _fill_text(line, ctx)).font = Font(name=font_name, size=size, color="595959")
        row += 1
    if row > 1:
        row += 1

    header_row = row
    header_fill = PatternFill("solid", fgColor=style.get("header_fill", "1F3864"))
    for index, col in enumerate(columns, 1):
        cell = sheet.cell(header_row, index, col["header"])
        cell.font = Font(name=font_name, size=size, bold=True, color=style.get("header_font", "FFFFFF"))
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = border
        sheet.column_dimensions[get_column_letter(index)].width = col.get("width", 16)
    sheet.row_dimensions[header_row].height = 22

    zebra = PatternFill("solid", fgColor=style["zebra"]) if style.get("zebra") else None
    data = build_rows(events, preset, ctx)
    for offset, values in enumerate(data):
        r = header_row + 1 + offset
        for index, value in enumerate(values, 1):
            cell = sheet.cell(r, index, value)
            cell.font = Font(name=font_name, size=size)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if zebra and offset % 2 == 1:
                cell.fill = zebra

    last = header_row + len(data)
    if preset.get("totals") and data:
        total_frames = sum(e.get("duration_frames", 0) for e in events)
        r = last + 2
        sheet.cell(r, 1, preset["totals"]).font = Font(name=font_name, size=size, bold=True)
        sheet.cell(r, 2, "{}  ({} cues)".format(format_duration(total_frames, ctx["fps"], "tc"), len(events))).font = \
            Font(name=font_name, size=size, bold=True)

    if style.get("freeze", True):
        sheet.freeze_panes = sheet.cell(header_row + 1, 1)
    if style.get("autofilter", True) and data:
        sheet.auto_filter.ref = "A{}:{}{}".format(header_row, get_column_letter(len(columns)), last)
    if style.get("landscape", True):
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.print_title_rows = "{0}:{0}".format(header_row)
    workbook.save(output_path)
    return output_path


def write_into_template(events, preset, ctx, template_path, output_path):
    """Fill a client's own .xlsx: find its header row by matching preset headers (and aliases),
    then write one row per event underneath, copying the template's first body-row style.
    {project}-style tokens anywhere in the template are replaced too.

    Returns the list of template headers that matched a preset column.
    """
    from copy import copy
    workbook = openpyxl.load_workbook(template_path)
    sheet = workbook.active
    names = {}
    for index, col in enumerate(preset["columns"]):
        for label in [col["header"]] + list(col.get("aliases", [])):
            names.setdefault(key_name(label), index)

    best_row, best_map = None, {}
    for r in range(1, min(sheet.max_row, 60) + 1):
        mapping = {}
        for c in range(1, sheet.max_column + 1):
            value = sheet.cell(r, c).value
            if isinstance(value, str) and key_name(value) in names:
                mapping[c] = names[key_name(value)]
        if len(mapping) > len(best_map):
            best_row, best_map = r, mapping
    if best_row is None or len(best_map) < 2:
        raise ValueError("Couldn't find a header row in the template. It needs at least two column headings "
                         "that match the preset (e.g. {}).".format(", ".join(c["header"] for c in preset["columns"][:4])))

    for row_cells in sheet.iter_rows():
        for cell in row_cells:
            if isinstance(cell.value, str) and "{" in cell.value:
                cell.value = _fill_text(cell.value, ctx)

    style_row = best_row + 1
    styles = {c: sheet.cell(style_row, c) for c in range(1, sheet.max_column + 1)}
    snapshots = {c: (copy(s.font), copy(s.fill), copy(s.border), copy(s.alignment), s.number_format)
                 for c, s in styles.items()}
    data = build_rows(events, preset, ctx)
    for offset, values in enumerate(data):
        r = best_row + 1 + offset
        for c in range(1, sheet.max_column + 1):
            cell = sheet.cell(r, c)
            font, fill, border, alignment, number_format = snapshots[c]
            cell.font, cell.fill, cell.border, cell.alignment, cell.number_format = (
                copy(font), copy(fill), copy(border), copy(alignment), number_format)
            cell.value = values[best_map[c]] if c in best_map else None
    workbook.save(output_path)
    return [sheet.cell(best_row, c).value for c in sorted(best_map)]


def build_report(source_path, output_path, preset, *, fps=DEFAULT_FPS, project="", track_filter="",
                 metadata_path=None, key_column=None, event_field="clip_name", merge=None,
                 template_path=None, include_gaps=False, fields=None, break_seconds=None,
                 df_deduct_seconds=None):
    """Run the whole pipeline and return a summary dict for the UI log."""
    events, meta = parse_file(source_path, fps)
    fps = meta["fps"]
    summary = {"parsed": len(events), "unmatched": [], "unused": [], "matched": None, "template_columns": None}
    if preset.get("layout") == "segments":
        if template_path:
            raise ValueError("Client templates can't be used with the segment timing sheet yet.")
        events = filter_tracks(events, track_filter)
        seconds = break_seconds if break_seconds is not None else preset.get("break_seconds")
        low, high = 1, None
        if seconds:
            low = max(1, int(round((float(seconds) - 0.5) * fps)))
            high = int(round((float(seconds) + 0.5) * fps))
        segments = build_segments(events, fps, low, high)
        if not segments:
            raise ValueError("No program found. Check the track filter and that the EDL has record timecodes.")
        ctx = make_context(project, source_path, meta, fields)
        ctx["df_deduct_seconds"] = df_deduct_seconds if df_deduct_seconds is not None else preset.get("df_deduct_seconds", 0)
        write_segment_sheet(segments, preset, ctx, output_path)
        program = [s for s in segments if s["kind"] == "segment"]
        summary.update(written=len(segments), segments=len(program), breaks=len(segments) - len(program),
                       runtime_frames=sum(s["frames"] for s in program), fps=fps)
        return summary
    if not include_gaps:
        events = drop_gaps(events)
    events = filter_tracks(events, track_filter)
    if merge is None:
        merge = preset.get("merge_back_to_back", False)
    if merge:
        events = merge_back_to_back(collapse_duplicates(events), fps)
    if metadata_path:
        headers, records = read_table(metadata_path)
        if not headers:
            raise ValueError("The metadata file is empty.")
        key = key_column if key_column in headers else headers[0]
        summary["matched"], summary["unmatched"], summary["unused"] = merge_metadata(
            events, records, key, event_field)
    sort_key = preset.get("sort", "rec_in")
    if sort_key == "rec_in":
        events.sort(key=lambda e: (tc_to_frames(e["rec_in"], fps) if e.get("rec_in") else 0, str(e.get("track", ""))))
    elif sort_key:
        events.sort(key=lambda e: str(e.get(key_name(sort_key), "")).casefold())
    ctx = make_context(project, source_path, meta, fields)
    if template_path:
        summary["template_columns"] = write_into_template(events, preset, ctx, template_path, output_path)
    else:
        write_xlsx(events, preset, ctx, output_path)
    summary["written"] = len(events)
    summary["runtime_frames"] = sum(e.get("duration_frames", 0) for e in events)
    summary["fps"] = fps
    return summary
