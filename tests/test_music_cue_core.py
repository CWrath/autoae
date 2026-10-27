import datetime
import os
import tempfile
import unittest

import docx

import edl_report_core as core
import music_cue_core as mc

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(HERE, "assets", "templates", "paramount_cue_sheet.docx")

EDL = """TITLE: Ep 101
FCM: NON-DROP FRAME

001  AX       V     C        00:00:00:00 00:01:00:00 01:00:00:00 01:01:00:00
002  AX       A13   C        00:00:00:00 00:00:52:00 01:00:28:19 01:01:20:03
* FROM CLIP NAME: CRYING_GAME_full.wav
003  AX       A14   C        00:00:00:00 00:00:52:00 01:00:28:19 01:01:20:03
* FROM CLIP NAME: CRYING_GAME_full.wav
004  AX       A13   C        00:00:00:00 00:00:10:00 01:02:00:00 01:02:10:00
* FROM CLIP NAME: OTHER_SONG.wav
005  AX       A13   C        00:00:10:00 00:00:20:00 01:02:10:00 01:02:20:00
* FROM CLIP NAME: OTHER_SONG.wav
006  AX       A5    C        00:00:00:00 00:00:05:00 01:03:00:00 01:03:05:00
* FROM CLIP NAME: DIALOGUE.wav
"""
LIBRARY = ("File,Title,Composer,Publisher,PRO,Label\n"
           "CRYING_GAME_full,CRYING GAME - Full Version,A (50%),Pub A,PRS,Label X\n")


class MusicCueTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.edl = os.path.join(self.tmp.name, "a.edl")
        self.lib = os.path.join(self.tmp.name, "lib.csv")
        self.out = os.path.join(self.tmp.name, "out.docx")
        open(self.edl, "w").write(EDL)
        open(self.lib, "w").write(LIBRARY)

    def build(self, **kw):
        logs = []
        count = mc.build_cue_sheet(self.edl, TEMPLATE, self.out, fps=24, track_filter="A13-A17",
                                   library_path=self.lib, key_column="File", log=logs.append,
                                   now=datetime.datetime(2026, 10, 22, 16, 28), **kw)
        return count, logs

    def test_time_cell(self):
        self.assertEqual(mc.time_cell("01:00:28:19", "01:01:20:03"), "01:00:28/01:01:20 (00:00:52)")
        self.assertEqual(mc.time_cell("01:00:00:00", "01:00:00:10"), "01:00:00/01:00:00 (00:00:01)")

    def test_fills_template(self):
        count, logs = self.build(info={"Show Title": "Show", "Cart #": "X1"}, sheet_id="42")
        self.assertEqual(count, 2)
        d = docx.Document(self.out)
        info = {}
        for row in d.tables[0].rows:
            cells = row.cells
            info[cells[0].text.strip()] = cells[1].text
        self.assertEqual(info["Show Title:"] if "Show Title:" in info else info["Show Title"], "Show")
        rows = d.tables[1].rows
        self.assertEqual(len(rows), 3)
        first = [c.text for c in rows[1].cells]
        self.assertEqual(first[0], "1")
        self.assertEqual(first[2], "CRYING GAME - Full Version")
        self.assertEqual(first[4], "A (50%)")
        self.assertEqual(first[6], "01:00:28/01:01:20 (00:00:52)")
        self.assertEqual(first[7], "BI")
        self.assertEqual(first[9], "PRS")
        second = [c.text for c in rows[2].cells]
        self.assertEqual(second[2], "OTHER SONG")  # fallback, back-to-back cuts merged
        self.assertEqual(second[6], "01:02:00/01:02:20 (00:00:20)")
        self.assertTrue(any("no library entry for 'OTHER SONG'" in m for m in logs))
        text = [p.text for p in d.paragraphs]
        self.assertIn("Usage Type: BI=Background Instrumental", text)
        self.assertTrue(any(t.startswith("*Report Generated on 10/22/2026 4:28 PM") and t.endswith("Cue sheet ID: 42")
                            for t in text))

    def test_song_key_groups_stems(self):
        key = mc.song_key
        self.assertEqual(key("XRC088_08_10 THE PAST JUDGEMENT_STRINGS.WAV"), key("XRC088_08 THE PAST JUDGEMENT.WAV"))
        self.assertEqual(key("C3 SIG_EQUAL_FULL MIX_SJJ JV.WAV"), key("C3 SIG_EQUAL_MIX ONLY PIANOS_SJJ JV.WAV"))
        self.assertEqual(key("CRUSH_102_INVESTIGATIONS - V2 -FULLMIX.WAV"), key("CRUSH_102_INVESTIGATIONS - V2 - SLOWARPSTEM.WAV"))
        self.assertNotEqual(key("C3 SIG_EQUAL_FULL MIX.WAV"), key("C3 SIG_ORGANISM_FULL MIX.WAV"))

    def test_stems_merge_but_separate_uses_do_not(self):
        edl = ("001  A  A15  C  00:00:00:00 00:00:10:00 01:00:00:00 01:00:10:00\n* FROM CLIP NAME: SONG_FULL MIX.WAV\n"
               "002  A  A16  C  00:00:00:00 00:00:08:00 01:00:02:00 01:00:10:00\n* FROM CLIP NAME: SONG_MIX ONLY PIANO.WAV\n"
               "003  A  A15  C  00:00:00:00 00:00:10:00 01:05:00:00 01:05:10:00\n* FROM CLIP NAME: SONG_FULL MIX.WAV\n")
        open(self.edl, "w").write(edl)
        cues, _ = mc.build_cues(core.parse_file(self.edl, 24)[0], 24, "A15-A16")
        self.assertEqual([(c["rec_in"], c["rec_out"]) for c in cues],
                         [("01:00:00:00", "01:00:10:00"), ("01:05:00:00", "01:05:10:00")])

    def test_exclude(self):
        cues, _ = mc.build_cues(core.parse_file(self.edl, 24)[0], 24, "A13-A17", exclude="other")
        self.assertEqual(len(cues), 1)

    PAGE = ("Police Car Radio Int\nLength: 1:02 Release Date: 04/06/2017 [Find Similar](https://x/y)\n\n"
            "### Filename:\nOTHER_SONG\n\n### Artist:\n[Valentino SFX / BMGPM](https://x/a)\n\n"
            "### Label:\n[Valentino Vintage SFX](https://x/l)\n\n### Publishers:\n"
            "NameRolePROIPIOwnership ShareCollection Share[BMG Production Works](https://x/p)OPBMI810363279100%\n\n"
            "### Writers:\nNameRolePROIPIShare[Robert Valentino](https://x/w)CBMI65539164100%[Jo Ann](https://x/w2)APRS12345678950%\n")

    def test_parse_track_page(self):
        record = mc.parse_track_pages(self.PAGE)[0]
        self.assertEqual(record["Filename"], "OTHER_SONG")
        self.assertEqual(record["Musical Composition"], "Police Car Radio Int")
        self.assertEqual(record["Performer(s)"], "Valentino SFX / BMGPM")
        self.assertEqual(record["Writer(s)"], "Robert Valentino (100%)-BMI / Jo Ann (50%)-PRS")
        self.assertEqual(record["Publisher(s)"], "BMG Production Works (100%)-BMI")
        self.assertEqual(record["Performing Rights Society"], "BMI / PRS")
        self.assertEqual(record["Record Label(s)"], "Valentino Vintage SFX")

    def test_parse_paramount_copy_without_filename_or_ipi(self):
        page = """Tommy Walter.Nervous Energy (No Drums&Perc)

### Artist:
[Tommy Walter](https://www.paramountglobalmusiclibrary.com/artist/35030)

### Album:
[Paramount Originals Vol. 1](https://www.paramountglobalmusiclibrary.com/album/12453882)

### Library:
[Paramount Originals](https://www.paramountglobalmusiclibrary.com/catalog/4814?page=labels)

### Label:
[Paramount Originals](https://www.paramountglobalmusiclibrary.com/label/622520?page=albums)

### BPM:
99

### Paramount Owned Cue Type:
[Music](https://www.paramountglobalmusiclibrary.com/track/15216497?s=Music)

### Has Vocals:
No

### Publishers:
NameRolePROIPIOwnership ShareCollection Share[Music By Video Inc.](https://www.paramountglobalmusiclibrary.com/publisher/849034)OP100%

### Writers:
NameRolePROIPIShare[Thomas Edward Walter](https://www.paramountglobalmusiclibrary.com/writer/849629)C50%[Andrew John Meyer](https://www.paramountglobalmusiclibrary.com/writer/849630)C50%
"""
        record = mc.parse_track_pages(page)[0]
        self.assertEqual(record["Filename"], "Tommy Walter.Nervous Energy (No Drums&Perc)")
        self.assertEqual(record["Musical Composition"], "Nervous Energy (No Drums&Perc)")
        self.assertEqual(record["Performer(s)"], "Tommy Walter")
        self.assertEqual(record["Album"], "Paramount Originals Vol. 1")
        self.assertEqual(record["Library"], "Paramount Originals")
        self.assertEqual(record["Record Label(s)"], "Paramount Originals")
        self.assertEqual(record["Publisher(s)"], "Music By Video Inc. (100%)")
        self.assertEqual(record["Writer(s)"],
                         "Thomas Edward Walter (50%) / Andrew John Meyer (50%)")
        self.assertEqual(record["Performing Rights Society"], "")
        self.assertEqual(mc.song_key(record["Filename"]),
                         mc.song_key("TOMMY WALTER_NERVOUS ENERGY (NO DRUMS&PERC).WAV"))

    def test_saved_tracks_fill_cues(self):
        cache = os.path.join(self.tmp.name, "tracks.json")
        mc.add_to_cache(cache, mc.parse_track_pages(self.PAGE))
        missing = []
        mc.build_cue_sheet(self.edl, TEMPLATE, self.out, fps=24, track_filter="A13-A17",
                           cache_records=mc.load_cache(cache), missing=missing, log=lambda m: None)
        rows = docx.Document(self.out).tables[1].rows
        other = [c.text for c in rows[2].cells]
        self.assertEqual(other[2], "Police Car Radio Int")
        self.assertEqual(other[5], "BMG Production Works (100%)-BMI")
        self.assertEqual(other[10], "Valentino Vintage SFX")
        self.assertEqual(missing, ["CRYING GAME FULL"])

    def test_no_music_raises(self):
        with self.assertRaises(ValueError):
            mc.build_cue_sheet(self.edl, TEMPLATE, self.out, fps=24, track_filter="A99")

    def test_library_template(self):
        out = os.path.join(self.tmp.name, "lib.xlsx")
        self.assertEqual(mc.write_library_template(self.edl, out, 24, "A13-A14"), 2)


if __name__ == "__main__":
    unittest.main()
