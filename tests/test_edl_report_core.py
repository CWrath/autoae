import tempfile
import unittest
from pathlib import Path

import openpyxl

import edl_report_core as core

EDL = """TITLE: Episode 101
FCM: NON-DROP FRAME

001  MUSIC1   A     C        00:00:00:00 00:00:10:00 01:00:00:00 01:00:10:00
* FROM CLIP NAME: Sunrise Theme.wav
* COMMENT: main title
002  MUSIC1   A     C        00:00:10:00 00:00:20:00 01:00:10:00 01:00:20:00
* FROM CLIP NAME: Sunrise Theme.wav
003  BL       V     C        00:00:00:00 00:00:05:00 01:00:20:00 01:00:25:00
004  SFX9     A2    C        00:00:00:00 00:00:02:12 01:00:30:00 01:00:32:12
* FROM CLIP NAME: Door Slam.wav
* SOURCE FILE: C:\\audio\\door.wav
"""


class TimecodeTests(unittest.TestCase):
    def test_round_trip_non_drop(self):
        frames = core.tc_to_frames("01:02:03:04", 24)
        self.assertEqual(frames, ((1 * 60 + 2) * 60 + 3) * 24 + 4)
        self.assertEqual(core.frames_to_tc(frames, 24), "01:02:03:04")

    def test_drop_frame_round_trip(self):
        for tc in ("00:01:00;02", "00:09:59;29", "00:10:00;00", "01:00:00;00"):
            frames = core.tc_to_frames(tc, 29.97)
            self.assertEqual(core.frames_to_tc(frames, 29.97, drop=True), tc)

    def test_drop_frame_known_value(self):
        self.assertEqual(core.tc_to_frames("00:01:00;02", 29.97), 1800)
        self.assertEqual(core.tc_to_frames("01:00:00;00", 29.97), 107892)


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def write(self, name, text):
        path = self.root / name
        path.write_text(text, encoding="utf-8")
        return path

    def test_edl_events_comments_and_duration(self):
        events, meta = core.parse_edl(self.write("a.edl", EDL), 24)
        self.assertEqual(meta["title"], "Episode 101")
        self.assertEqual(len(events), 4)
        self.assertEqual(events[0]["clip_name"], "Sunrise Theme.wav")
        self.assertEqual(events[0]["duration_frames"], 240)
        self.assertEqual(events[0]["comment"], "main title")
        self.assertEqual(events[3]["source_file"], "C:\\audio\\door.wav")
        self.assertEqual(events[3]["track"], "A2")

    def test_edl_without_clip_name_falls_back_to_reel(self):
        edl = "001  TAPE01  V  C  00:00:00:00 00:00:01:00 01:00:00:00 01:00:01:00\n"
        events, _ = core.parse_edl(self.write("b.edl", edl), 24)
        self.assertEqual(events[0]["clip_name"], "TAPE01")

    def test_ale(self):
        ale = ("Heading\nFIELD_DELIM\tTABS\nFPS\t25\n\nColumn\nName\tTape\tStart\tEnd\tLicensor\n\n"
               "Data\nCue A\tT1\t01:00:00:00\t01:00:05:00\tAcme\n")
        events, meta = core.parse_ale(self.write("c.ale", ale))
        self.assertEqual(meta["fps"], 25.0)
        self.assertEqual(events[0]["clip_name"], "Cue A")
        self.assertEqual(events[0]["licensor"], "Acme")
        self.assertEqual(events[0]["duration_frames"], 125)

    def test_xmeml(self):
        xml = """<xmeml version="4"><sequence><name>Seq</name><rate><timebase>24</timebase><ntsc>FALSE</ntsc></rate>
        <timecode><string>01:00:00:00</string><displayformat>NDF</displayformat></timecode>
        <media><audio><track><clipitem><name>Song.wav</name><start>0</start><end>48</end><in>0</in><out>48</out>
        <file id="f1"><name>Song.wav</name><pathurl>file:///a/Song.wav</pathurl></file></clipitem></track></audio></media>
        </sequence></xmeml>"""
        events, meta = core.parse_xml(self.write("d.xml", xml))
        self.assertEqual(meta["title"], "Seq")
        self.assertEqual(events[0]["rec_in"], "01:00:00:00")
        self.assertEqual(events[0]["duration_frames"], 48)
        self.assertEqual(events[0]["track"], "A1")


class ProcessingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.edl = self.root / "a.edl"
        self.edl.write_text(EDL, encoding="utf-8")

    def tearDown(self):
        self.temp.cleanup()

    def test_gap_filter_and_merge(self):
        events, _ = core.parse_edl(self.edl, 24)
        events = core.drop_gaps(events)
        self.assertEqual(len(events), 3)
        merged = core.merge_back_to_back(events, 24)
        self.assertEqual(len(merged), 2)
        self.assertEqual(merged[0]["duration_frames"], 480)
        self.assertEqual(merged[0]["rec_out"], "01:00:20:00")

    def test_track_filter(self):
        events, _ = core.parse_edl(self.edl, 24)
        self.assertEqual(len(core.filter_tracks(events, "A2")), 1)
        self.assertEqual(len(core.filter_tracks(events, "A*")), 3)
        self.assertEqual(len(core.filter_tracks(events, "")), 4)

    def test_track_ranges(self):
        self.assertEqual(core.expand_track_patterns("A13-A15, V*, a2"), ["a13", "a14", "a15", "v*", "a2"])
        self.assertEqual(core.expand_track_patterns("A3-5"), ["a3", "a4", "a5"])
        self.assertEqual(core.expand_track_patterns("A9-A3"), ["a9-a3"])

    def test_metadata_merge_ignores_extension_and_reports_unmatched(self):
        events, _ = core.parse_edl(self.edl, 24)
        events = core.drop_gaps(events)
        meta = self.root / "meta.csv"
        meta.write_text("Title,Licensor,Status\nsunrise theme,Acme,Cleared\nUnused,X,Y\n", encoding="utf-8")
        headers, records = core.read_table(meta)
        matched, unmatched, unused = core.merge_metadata(events, records, headers[0])
        self.assertEqual(matched, 2)
        self.assertEqual(unmatched, ["Door Slam.wav"])
        self.assertEqual(unused, ["unused"])
        self.assertEqual(events[0]["licensor"], "Acme")

    def test_build_report_writes_styled_xlsx(self):
        meta = self.root / "meta.csv"
        meta.write_text("Title,Licensor,Status\nSunrise Theme,Acme,Cleared\n", encoding="utf-8")
        out = self.root / "out.xlsx"
        summary = core.build_report(self.edl, out, core.BUILTIN_PRESETS["Clearance Report"], fps=24,
                                    project="Show", metadata_path=meta, key_column="Title")
        self.assertEqual(summary["written"], 2)
        self.assertEqual(summary["unmatched"], ["Door Slam.wav"])
        sheet = openpyxl.load_workbook(out).active
        self.assertEqual(sheet["A1"].value, "Show  |  Clearance Report")
        rows = [[c.value for c in r] for r in sheet.iter_rows(min_row=1)]
        header = next(r for r in rows if r[0] == "#")
        first = rows[rows.index(header) + 1]
        self.assertEqual(first[1], "Sunrise Theme.wav")
        self.assertEqual(first[3], "01:00:00:00")
        self.assertEqual(first[5], "00:00:20:00")
        self.assertEqual(first[7], "Acme")
        self.assertEqual(first[10], "Cleared")

    def test_template_is_filled_by_header_names(self):
        template = self.root / "client.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet["A1"] = "Production: {project}"
        for col, name in enumerate(["Cue Title", "Start", "End", "Length", "Notes"], 1):
            sheet.cell(4, col, name)
        sheet["A5"].font = openpyxl.styles.Font(bold=True, color="FF0000")
        book.save(template)
        out = self.root / "filled.xlsx"
        summary = core.build_report(self.edl, out, core.BUILTIN_PRESETS["Clearance Report"], fps=24,
                                    project="My Show", template_path=template)
        self.assertEqual(len(summary["template_columns"]), 5)
        result = openpyxl.load_workbook(out).active
        self.assertEqual(result["A1"].value, "Production: My Show")
        self.assertEqual(result["A5"].value, "Sunrise Theme.wav")
        self.assertEqual(result["B5"].value, "01:00:00:00")
        self.assertTrue(result["A6"].font.bold)

    def test_template_without_matching_headers_fails_clearly(self):
        template = self.root / "bad.xlsx"
        book = openpyxl.Workbook()
        book.active["A1"] = "nothing useful"
        book.save(template)
        with self.assertRaises(ValueError):
            core.build_report(self.edl, self.root / "x.xlsx", core.BUILTIN_PRESETS["Clearance Report"],
                              fps=24, template_path=template)


class SegmentTimingTests(unittest.TestCase):
    # Segment / break layout from a delivered timing sheet (23.98), as record in -> out (exclusive).
    CUT = [("01:00:00:00", "01:07:49:13", "S1"), ("01:07:49:13", "01:07:51:13", "BL"),
           ("01:07:51:13", "01:13:48:05", "S2"), ("01:13:48:05", "01:13:50:05", "BL"),
           ("01:13:50:05", "01:20:05:19", "S3"), ("01:20:05:19", "01:20:07:19", "BL"),
           ("01:20:07:19", "01:24:37:13", "S4"), ("01:24:37:13", "01:24:39:13", "BL"),
           ("01:24:39:13", "01:30:26:05", "S5"), ("01:30:26:05", "01:30:28:05", "BL"),
           ("01:30:28:05", "01:42:37:11", "S6")]

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        lines = ["TITLE: Ep"]
        for number, (a, b, reel) in enumerate(self.CUT, 1):
            lines.append("{:03d}  {:<8} V     C        00:00:00:00 00:00:01:00 {} {}".format(number, reel, a, b))
        # audio running under the breaks must not hide them
        lines.append("{:03d}  MUSIC   A     C        00:00:00:00 00:42:37:00 01:00:00:00 01:42:37:00".format(len(self.CUT) + 1))
        self.edl = self.root / "cut.edl"
        self.edl.write_text("\n".join(lines) + "\n", encoding="utf-8")

    def tearDown(self):
        self.temp.cleanup()

    def test_segments_and_breaks_match_the_sheet(self):
        events, _ = core.parse_edl(self.edl, 23.976)
        segments = core.build_segments(events, 23.976)
        self.assertEqual([s["label"] for s in segments][:4],
                         ["SEGMENT 1", "Break 1 (Black)", "SEGMENT 2", "Break 2 (Black)"])
        self.assertEqual(len([s for s in segments if s["kind"] == "segment"]), 6)
        self.assertEqual(segments[0]["frames"], 11269)
        self.assertEqual(segments[1]["frames"], 48)

    def test_df_conversion(self):
        self.assertEqual(core.to_df_timecode(11269, 23.976), "00;07;50;00")
        self.assertEqual(core.to_df_timecode(8560, 23.976), "00;05;57;00")
        self.assertEqual(core.to_df_timecode(6474, 23.976), "00;04;30;01")
        self.assertEqual(core.to_df_timecode(48, 23.976), "00;00;02;00")
        self.assertEqual(core.round_df_timecode("00;06;15;29"), "00;06;16;00")
        self.assertEqual(core.round_df_timecode("00;04;30;01"), "00;04;30;00")

    def test_workbook_layout(self):
        out = self.root / "t.xlsx"
        summary = core.build_report(self.edl, out, core.BUILTIN_PRESETS["Segment Timing Sheet"], fps=23.976,
                                    fields={"series": "Final Twist", "episode_title": "A Recipe",
                                            "episode_number": "204"})
        self.assertEqual((summary["segments"], summary["breaks"]), (6, 5))
        sheet = openpyxl.load_workbook(out).active
        self.assertEqual(sheet["A1"].value, "Series Title: Final Twist")
        self.assertEqual(sheet["A4"].value, "Program ID: N/A")
        self.assertEqual([sheet.cell(6, c).value for c in (1, 2, 4, 5, 6)],
                         ["23.98 Timecode In", "23.98 Timecode Out", "Length", "DF Conversion", "DF Rounded"])
        self.assertEqual([sheet.cell(7, c).value for c in range(1, 7)],
                         ["01:00:00:00", "01:07:49:12", "SEGMENT 1", "00:07:49:13", "00;07;50;00", "00;07;50;00"])
        self.assertEqual([sheet.cell(8, c).value for c in range(1, 6)],
                         ["01:07:49:13", "01:07:51:12", "Break 1 (Black)", "00:00:02:00", "00;00;02;00"])
        self.assertIsNone(sheet["F8"].value)
        self.assertIsNone(sheet["A9"].value)
        self.assertEqual(sheet["A10"].value, "01:07:51:13")
        self.assertEqual(sheet["C10"].value, "SEGMENT 2")
        self.assertEqual(sheet["B22"].value, "01:42:37:10")
        self.assertEqual(sheet["C24"].value, "TPT in 23.98")
        self.assertEqual(sheet["C7"].font.name, "Arial")

    def test_gap_between_events_becomes_a_break(self):
        edl = ("001  A  V  C  00:00:00:00 00:00:05:00 01:00:00:00 01:00:05:00\n"
               "002  B  V  C  00:00:00:00 00:00:05:00 01:00:07:00 01:00:12:00\n")
        path = self.root / "gap.edl"
        path.write_text(edl, encoding="utf-8")
        events, _ = core.parse_edl(path, 24)
        segments = core.build_segments(events, 24)
        self.assertEqual([s["label"] for s in segments], ["SEGMENT 1", "Break 1", "SEGMENT 2"])
        self.assertEqual(segments[1]["frames"], 48)

    def test_gap_outside_break_length_stays_in_the_segment(self):
        edl = ("001  A  V  C  00:00:00:00 00:00:10:00 01:00:00:00 01:00:10:00\n"
               "002  B  V  C  00:00:00:00 00:00:10:00 01:00:12:00 01:00:22:00\n"
               "003  C  V  C  00:00:00:00 00:00:10:00 01:00:25:00 01:00:35:00\n")
        path = self.root / "gaps.edl"
        path.write_text(edl, encoding="utf-8")
        out = self.root / "g.xlsx"
        summary = core.build_report(path, out, core.BUILTIN_PRESETS["Segment Timing Sheet"], fps=24)
        self.assertEqual((summary["segments"], summary["breaks"]), (2, 1))
        sheet = openpyxl.load_workbook(out).active
        self.assertEqual(sheet["C8"].value, "Break 1 (Black)")
        self.assertEqual(sheet["B10"].value, "01:00:34:23")
        self.assertEqual(sheet["D10"].value, "00:00:23:00")
        self.assertEqual(sheet["D12"].value, "00:00:34:23")  # TPT counts first in to last out
        self.assertEqual(sheet["F12"].value, "00:00:30:29")  # DF total: program only, less the 2 s card

    def test_clip_name_uses_source_file_when_reel_is_truncated(self):
        edl = ("001  VERY_LONG_REEL_NAME_CUT_AT_32_CH  A  C  00:00:00:00 00:00:01:00 01:00:00:00 01:00:01:00\n"
               "*SOURCE FILE: Very Long Reel Name Cut At 32 Characters.WAV\n")
        path = self.root / "avid.edl"
        path.write_text(edl, encoding="utf-8")
        events, _ = core.parse_edl(path, 24)
        self.assertEqual(events[0]["clip_name"], "Very Long Reel Name Cut At 32 Characters.WAV")

    def test_stereo_pairs_collapse_to_one_line(self):
        edl = ("001  SONG  A   C  00:00:00:00 00:00:05:00 01:00:00:00 01:00:05:00\n"
               "002  SONG  A2  C  00:00:00:00 00:00:05:00 01:00:00:00 01:00:05:00\n")
        path = self.root / "pair.edl"
        path.write_text(edl, encoding="utf-8")
        events, _ = core.parse_edl(path, 24)
        collapsed = core.collapse_duplicates(events)
        self.assertEqual(len(collapsed), 1)
        self.assertEqual(collapsed[0]["track"], "A/A2")

    def test_template_rejected_for_segment_layout(self):
        with self.assertRaises(ValueError):
            core.build_report(self.edl, self.root / "x.xlsx", core.BUILTIN_PRESETS["Segment Timing Sheet"],
                              fps=23.976, template_path=self.root / "t.xlsx")


if __name__ == "__main__":
    unittest.main()
